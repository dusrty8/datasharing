#!/usr/bin/env python3
"""Builds the Cortex Atlas ONE-PAGER (single portrait Letter page)."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

INK=RGBColor(0x09,0x13,0x1B); INK2=RGBColor(0x0E,0x1A,0x22); ORANGE=RGBColor(0xF8,0x97,0x39)
SURFACE=RGBColor(0xF4,0xF4,0xF4); CARD=RGBColor(0xFF,0xFF,0xFF); LINE=RGBColor(0xD6,0xDC,0xE4)
TEXT2=RGBColor(0x59,0x59,0x59); MUTE=RGBColor(0xA9,0xAF,0xB6); GREEN=RGBColor(0x1E,0x7F,0x4F)
RED=RGBColor(0xE8,0x45,0x1C); BLUE=RGBColor(0x05,0x63,0xC1); WHITE=RGBColor(0xFF,0xFF,0xFF)
ORANGE_D=RGBColor(0x9a,0x5c,0x12); STROKE=RGBColor(0x2a,0x3a,0x44)
FONT="Calibri"

prs=Presentation(); prs.slide_width=Inches(8.5); prs.slide_height=Inches(11)
SW,SH=prs.slide_width,prs.slide_height; BLANK=prs.slide_layouts[6]
MX=Inches(0.5); CW=Inches(7.5)
s=prs.slides.add_slide(BLANK); s.background.fill.solid(); s.background.fill.fore_color.rgb=SURFACE

def no_shadow(shp):
    el=shp._element.spPr
    if el.find(qn('a:effectLst')) is None: el.append(el.makeelement(qn('a:effectLst'),{}))
def rect(x,y,w,h,fill=None,line=None,lw=0.75,shape=MSO_SHAPE.RECTANGLE,radius=None):
    shp=s.shapes.add_shape(shape,x,y,w,h)
    if fill is None: shp.fill.background()
    else: shp.fill.solid(); shp.fill.fore_color.rgb=fill
    if line is None: shp.line.fill.background()
    else: shp.line.color.rgb=line; shp.line.width=Pt(lw)
    no_shadow(shp)
    if radius is not None and shape==MSO_SHAPE.ROUNDED_RECTANGLE:
        try: shp.adjustments[0]=radius
        except Exception: pass
    return shp
def tb(x,y,w,h,anchor=MSO_ANCHOR.TOP):
    b=s.shapes.add_textbox(x,y,w,h); tf=b.text_frame; tf.word_wrap=True; tf.vertical_anchor=anchor
    tf.margin_left=Pt(0); tf.margin_right=Pt(0); tf.margin_top=Pt(0); tf.margin_bottom=Pt(0); return tf
def par(tf,text,size=10,bold=False,color=INK,align=PP_ALIGN.LEFT,before=0,after=3,name=FONT,first=False):
    p=tf.paragraphs[0] if first else tf.add_paragraph()
    p.text=text; p.alignment=align; p.space_before=Pt(before); p.space_after=Pt(after)
    for r in p.runs: r.font.size=Pt(size); r.font.bold=bold; r.font.color.rgb=color; r.font.name=name
    return p
def card(x,y,w,h,fill=CARD,line=LINE): return rect(x,y,w,h,fill=fill,line=line,lw=1.0,shape=MSO_SHAPE.ROUNDED_RECTANGLE,radius=0.06)
def kicker(x,y,text,color=ORANGE):
    rect(x,y+Inches(0.02),Inches(0.22),Inches(0.07),fill=color)
    t=tb(x+Inches(0.3),y-Inches(0.04),Inches(5.5),Inches(0.3)); par(t,text.upper(),10.5,True,color,first=True,after=0)
def mini(tf,lead,rest,size=9.5,first=False):
    p=tf.paragraphs[0] if first else tf.add_paragraph()
    p.space_before=Pt(0); p.space_after=Pt(4); p.alignment=PP_ALIGN.LEFT
    r0=p.add_run(); r0.text="— "; r0.font.size=Pt(size); r0.font.bold=True; r0.font.color.rgb=ORANGE; r0.font.name=FONT
    r1=p.add_run(); r1.text=lead+"  "; r1.font.size=Pt(size); r1.font.bold=True; r1.font.color.rgb=INK; r1.font.name=FONT
    r2=p.add_run(); r2.text=rest; r2.font.size=Pt(size); r2.font.color.rgb=TEXT2; r2.font.name=FONT

# ===== HEADER BAND =====
rect(Inches(0),Inches(0),SW,Inches(1.55),fill=INK)
rect(Inches(0),Inches(0),SW,Inches(0.12),fill=ORANGE)
t=tb(MX,Inches(0.32),Inches(5.4),Inches(0.7)); par(t,"Cortex Atlas",30,True,WHITE,first=True,after=2)
t=tb(MX,Inches(0.95),Inches(5.4),Inches(0.4)); par(t,"Commercial data governance, connected.",13.5,False,RGBColor(0xC7,0xCC,0xD3),first=True,after=0)
rect(SW-MX-Inches(2.0),Inches(0.36),Inches(2.0),Inches(0.4),fill=INK2,line=STROKE,lw=1.0,shape=MSO_SHAPE.ROUNDED_RECTANGLE,radius=0.5)
t=tb(SW-MX-Inches(2.0),Inches(0.36),Inches(2.0),Inches(0.4),anchor=MSO_ANCHOR.MIDDLE); par(t,"EXECUTIVE ONE-PAGER",9.5,True,ORANGE,align=PP_ALIGN.CENTER,first=True,after=0)
t=tb(SW-MX-Inches(2.0),Inches(0.9),Inches(2.0),Inches(0.5)); par(t,"For Commercial Insights &\nData Procurement",9,False,MUTE,align=PP_ALIGN.RIGHT,first=True,after=0)

# ===== CHALLENGE + ESTATE =====
y=Inches(1.78)
kicker(MX,y,"The challenge & the idea")
t=tb(MX,y+Inches(0.32),CW,Inches(0.5))
par(t,"Five tools — requests, catalogue, KPIs, vendors, agreements — run today as disconnected SharePoint forms. Cortex Atlas connects them into one estate: ",9.8,False,TEXT2,first=True,after=0)
t2=tb(MX,y+Inches(0.66),CW,Inches(0.3)); par(t2,'"One asset, five tools" — a request becomes a catalogued asset, measured by KPIs, supplied by a vendor, governed by an agreement, every link clickable.',9.8,True,INK,first=True,after=0)
# 5-tool strip
flow=[("DPRH","Data Requests"),("ASSET","Data Catalogue"),("METRIC","KPI Catalogue"),("VENDOR","Vendor Hub"),("TPA","TPA Hub")]
sy=y+Inches(1.05); cw=Inches(1.42)
for i,(tag,nm) in enumerate(flow):
    xx=MX+i*(cw+Inches(0.1)); card(xx,sy,cw,Inches(0.72))
    tt=tb(xx+Inches(0.12),sy+Inches(0.1),cw-Inches(0.24),Inches(0.55)); par(tt,tag,8,True,ORANGE,first=True,after=2); par(tt,nm,10.5,True,INK,after=0)
    if i<4:
        ar=tb(xx+cw-Inches(0.01),sy+Inches(0.2),Inches(0.12),Inches(0.35),anchor=MSO_ANCHOR.MIDDLE); par(ar,"→",12,True,ORANGE,align=PP_ALIGN.CENTER,first=True,after=0)

# ===== TWO COLUMNS =====
colY=Inches(4.15); colH=Inches(4.35); LW=Inches(3.65); RW=Inches(3.65); RX=MX+LW+Inches(0.2)
# LEFT column
card(MX,colY,LW,colH)
lx=MX+Inches(0.22); lw=LW-Inches(0.44)
ty=colY+Inches(0.2)
t=tb(lx,ty,lw,Inches(0.3)); par(t,"WHAT IT DOES",10.5,True,ORANGE,first=True,after=0)
tt=tb(lx,ty+Inches(0.32),lw,Inches(1.9))
mini(tt,"Data Requests","status-driven, role-gated form; duplicate-buy detection (AI) at intake.",first=True)
mini(tt,"Data Catalogue","system of record; auto PII/PHI classification; quality flags.")
mini(tt,"KPI Catalogue","every metric with lineage back to its source asset.")
mini(tt,"Vendor Hub","what each vendor holds, purchased vs available.")
mini(tt,"TPA Hub","agreement register with a live 30/60/90 reminder engine.")
ty2=colY+Inches(2.5)
t=tb(lx,ty2,lw,Inches(0.3)); par(t,"THREE LAYERS",10.5,True,ORANGE,first=True,after=0)
tt=tb(lx,ty2+Inches(0.32),lw,Inches(1.4))
mini(tt,"Estate","the five tools, faithful to how they work today.",first=True)
mini(tt,"Automation","reminders, request↔catalogue sync, classification.")
mini(tt,"Intelligence","cross-tool signals, vendor value, dashboards & Ask Atlas (AI).")
# RIGHT column
card(RX,colY,RW,colH)
rx=RX+Inches(0.22); rw=RW-Inches(0.44)
ty=colY+Inches(0.2)
t=tb(rx,ty,rw,Inches(0.3)); par(t,"WHY IT MATTERS",10.5,True,GREEN,first=True,after=0)
tt=tb(rx,ty+Inches(0.32),rw,Inches(1.9))
mini(tt,"Stop duplicate spend","overlap caught before money is committed.",first=True)
mini(tt,"Never miss a renewal","automated, impact-aware expiry warnings.")
mini(tt,"Negotiate from strength","vendor value & concentration analytics.")
mini(tt,"Audit-ready by default","automatic classification and full lineage.")
mini(tt,"Decide with confidence","one source of truth + plain-English ask.")
ty2=colY+Inches(2.5)
t=tb(rx,ty2,rw,Inches(0.3)); par(t,"WHAT THE PROTOTYPE SURFACES",10.5,True,ORANGE,first=True,after=0)
tt=tb(rx,ty2+Inches(0.32),rw,Inches(1.6))
mini(tt,"~$1.09M","of duplicate-buy exposure flagged at intake.",first=True)
mini(tt,"6 agreements","expire within 90 days (3 within 30).")
mini(tt,"HHI ~1,825","top two vendors hold ~50% of spend.")
mini(tt,"1 Bronze asset","feeding a Global, strategic KPI.")

# ===== NEXT STEPS =====
ny=Inches(8.75)
kicker(MX,ny,"Next steps")
steps=[("1","Walk the prototype live"),("2","Baseline data audit"),("3","Confirm scope & architecture"),("4","Stand up Phase 1")]
cw=Inches(1.78); sy=ny+Inches(0.36)
for i,(n,tl) in enumerate(steps):
    xx=MX+i*(cw+Inches(0.06)); card(xx,sy,cw,Inches(0.72))
    rect(xx+Inches(0.14),sy+Inches(0.2),Inches(0.32),Inches(0.32),fill=ORANGE,shape=MSO_SHAPE.ROUNDED_RECTANGLE,radius=0.25)
    bt=tb(xx+Inches(0.14),sy+Inches(0.2),Inches(0.32),Inches(0.32),anchor=MSO_ANCHOR.MIDDLE); par(bt,n,11,True,INK,align=PP_ALIGN.CENTER,first=True,after=0)
    t=tb(xx+Inches(0.54),sy+Inches(0.1),cw-Inches(0.6),Inches(0.55),anchor=MSO_ANCHOR.MIDDLE); par(t,tl,9.8,True,INK,first=True,after=0)

# ===== FOOTER =====
rect(MX,Inches(9.95),CW,Pt(0.75),fill=LINE)
t=tb(MX,Inches(10.05),CW,Inches(0.6))
par(t,"Target stack: SharePoint · Power Automate · Microsoft Purview · Azure OpenAI.  Delivered in three phases — value from Phase 1.",9.5,True,INK,first=True,after=3)
par(t,"Illustrative prototype, synthetic data (NovaCura Pharmaceuticals). Figures directional pending a client baseline audit. All AI output is advisory — a human approves any action.",8.5,False,MUTE,after=0)

prs.save("/home/user/datasharing/Cortex-Atlas-One-Pager.pptx")
print("One-pager saved (portrait 8.5x11), shapes:", len(s.shapes))
