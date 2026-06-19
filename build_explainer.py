#!/usr/bin/env python3
"""Builds the Cortex Atlas USER EXPLAINER deck."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

INK=RGBColor(0x09,0x13,0x1B); INK2=RGBColor(0x0E,0x1A,0x22); ORANGE=RGBColor(0xF8,0x97,0x39)
SURFACE=RGBColor(0xF4,0xF4,0xF4); CARD=RGBColor(0xFF,0xFF,0xFF); LINE=RGBColor(0xD6,0xDC,0xE4)
TEXT2=RGBColor(0x59,0x59,0x59); MUTE=RGBColor(0xA9,0xAF,0xB6); GREEN=RGBColor(0x1E,0x7F,0x4F)
RED=RGBColor(0xE8,0x45,0x1C); BLUE=RGBColor(0x05,0x63,0xC1); WHITE=RGBColor(0xFF,0xFF,0xFF)
ORANGE_D=RGBColor(0x9a,0x5c,0x12); STROKE=RGBColor(0x2a,0x3a,0x44)
FONT="Calibri"

prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
SW,SH=prs.slide_width,prs.slide_height; BLANK=prs.slide_layouts[6]
MX=Inches(0.75); CW=Inches(13.333-1.5)

def slide(): return prs.slides.add_slide(BLANK)
def bg(s,c): s.background.fill.solid(); s.background.fill.fore_color.rgb=c
def no_shadow(shp):
    el=shp._element.spPr
    if el.find(qn('a:effectLst')) is None: el.append(el.makeelement(qn('a:effectLst'),{}))
def rect(s,x,y,w,h,fill=None,line=None,lw=0.75,shape=MSO_SHAPE.RECTANGLE,radius=None):
    shp=s.shapes.add_shape(shape,x,y,w,h)
    if fill is None: shp.fill.background()
    else: shp.fill.solid(); shp.fill.fore_color.rgb=fill
    if line is None: shp.line.fill.background()
    else: shp.line.color.rgb=line; shp.line.width=Pt(lw)
    no_shadow(shp)
    if radius is not None and shape==MSO_SHAPE.ROUNDED_RECTANGLE:
        try: shp.adjustments[0]=radius
        except Exception: pass
    return shp
def tb(s,x,y,w,h,anchor=MSO_ANCHOR.TOP):
    b=s.shapes.add_textbox(x,y,w,h); tf=b.text_frame; tf.word_wrap=True; tf.vertical_anchor=anchor
    tf.margin_left=Pt(0); tf.margin_right=Pt(0); tf.margin_top=Pt(0); tf.margin_bottom=Pt(0); return tf
def par(tf,text,size=14,bold=False,color=INK,align=PP_ALIGN.LEFT,before=0,after=4,name=FONT,first=False):
    p=tf.paragraphs[0] if first else tf.add_paragraph()
    p.text=text; p.alignment=align; p.space_before=Pt(before); p.space_after=Pt(after)
    for r in p.runs: r.font.size=Pt(size); r.font.bold=bold; r.font.color.rgb=color; r.font.name=name
    return p
def bullets(tf,items,size=13,color=TEXT2,after=7):
    first=True
    for it in items:
        lead,rest=it if isinstance(it,tuple) else (None,it)
        p=tf.paragraphs[0] if first else tf.add_paragraph(); first=False
        p.space_before=Pt(0); p.space_after=Pt(after); p.alignment=PP_ALIGN.LEFT
        r0=p.add_run(); r0.text="-  "; r0.font.size=Pt(size); r0.font.bold=True; r0.font.color.rgb=ORANGE; r0.font.name=FONT
        if lead:
            r1=p.add_run(); r1.text=lead+"  "; r1.font.size=Pt(size); r1.font.bold=True; r1.font.color.rgb=INK; r1.font.name=FONT
            r2=p.add_run(); r2.text=rest; r2.font.size=Pt(size); r2.font.color.rgb=color; r2.font.name=FONT
        else:
            r1=p.add_run(); r1.text=rest; r1.font.size=Pt(size); r1.font.color.rgb=color; r1.font.name=FONT
def badge(s,x,y,text,d=Inches(0.42),fill=ORANGE,fg=INK):
    rect(s,x,y,d,d,fill=fill,shape=MSO_SHAPE.ROUNDED_RECTANGLE,radius=0.25)
    t=tb(s,x,y,d,d,anchor=MSO_ANCHOR.MIDDLE); par(t,text,13.5,True,fg,align=PP_ALIGN.CENTER,first=True,after=0)
def card(s,x,y,w,h,fill=CARD,line=LINE): return rect(s,x,y,w,h,fill=fill,line=line,lw=1.0,shape=MSO_SHAPE.ROUNDED_RECTANGLE,radius=0.05)
def notes(s,t): s.notes_slide.notes_text_frame.text=t
def header(s,kicker,title,sub=None):
    rect(s,MX,Inches(0.62),Inches(0.34),Inches(0.09),fill=ORANGE)
    t=tb(s,MX+Inches(0.45),Inches(0.5),CW-Inches(0.45),Inches(0.35)); par(t,kicker.upper(),12.5,True,ORANGE,first=True,after=0)
    t2=tb(s,MX,Inches(0.92),CW,Inches(0.7)); par(t2,title,29,True,INK,first=True,after=0)
    if sub:
        t3=tb(s,MX,Inches(1.62),CW,Inches(0.55)); par(t3,sub,14.5,False,TEXT2,first=True,after=0)
def footer(s,n):
    rect(s,MX,Inches(7.02),CW,Pt(0.75),fill=LINE)
    t=tb(s,MX,Inches(7.08),Inches(9),Inches(0.3)); par(t,"Cortex Atlas  -  User explainer  -  Illustrative prototype, synthetic data",9,False,MUTE,first=True,after=0)
    t2=tb(s,SW-MX-Inches(1.0),Inches(7.08),Inches(1.0),Inches(0.3)); par(t2,str(n),9,False,MUTE,align=PP_ALIGN.RIGHT,first=True,after=0)
PAGE=0
def pg():
    global PAGE; PAGE+=1; return PAGE
def twocol(s,leftTitle,leftItems,rightTitle,rightItems,y=Inches(2.1),h=Inches(4.3)):
    card(s,MX,y,Inches(5.85),h); card(s,MX+Inches(6.15),y,Inches(5.85),h)
    lt=tb(s,MX+Inches(0.3),y+Inches(0.22),Inches(5.3),Inches(0.4)); par(lt,leftTitle,14.5,True,INK,first=True,after=0)
    li=tb(s,MX+Inches(0.3),y+Inches(0.7),Inches(5.3),h-Inches(0.9)); bullets(li,leftItems)
    rt=tb(s,MX+Inches(6.45),y+Inches(0.22),Inches(5.3),Inches(0.4)); par(rt,rightTitle,14.5,True,ORANGE_D,first=True,after=0)
    ri=tb(s,MX+Inches(6.45),y+Inches(0.7),Inches(5.3),h-Inches(0.9)); bullets(ri,rightItems)

# ---------- 1 Title
s=slide(); bg(s,INK)
rect(s,0,0,SW,Inches(0.16),fill=ORANGE)
rect(s,MX,Inches(1.5),Inches(2.4),Inches(0.5),fill=INK2,line=STROKE,lw=1.0,shape=MSO_SHAPE.ROUNDED_RECTANGLE,radius=0.5)
t=tb(s,MX,Inches(1.5),Inches(2.4),Inches(0.5),anchor=MSO_ANCHOR.MIDDLE); par(t,"USER EXPLAINER",11,True,ORANGE,align=PP_ALIGN.CENTER,first=True,after=0)
t=tb(s,MX,Inches(2.4),Inches(11.5),Inches(1.3)); par(t,"Cortex Atlas",58,True,WHITE,first=True,after=2)
t=tb(s,MX,Inches(3.6),Inches(11.5),Inches(0.8)); par(t,"Your connected workspace for commercial data governance.",24,False,RGBColor(0xC7,0xCC,0xD3),first=True,after=0)
t=tb(s,MX,Inches(4.6),Inches(11.5),Inches(1.0)); par(t,"What it is, how it is laid out, what each part shows, and what you can do in it. A companion to the full User Training Guide.",15,False,MUTE,first=True,after=0)
notes(s,"Orientation deck for the people who will actually use Atlas day to day - requesters, reviewers, catalogue stewards, vendor and access managers. Goal: by the end they understand what the tool is, how to move around it, and what each module does. Pair it with the User Training Guide for step-by-step instructions.")

# ---------- 2 What is it
s=slide(); bg(s,SURFACE)
header(s,"Start here","What is Cortex Atlas?")
t=tb(s,MX,Inches(2.0),CW,Inches(1.0))
par(t,"Cortex Atlas is one place to request, record, measure, source and govern the commercial data you buy.",18,True,INK,first=True,after=8)
par(t,"It brings together the five tools you use today - the purchase request hub, the data catalogue, the KPI catalogue, the vendor hub and the third-party-agreement register - into a single connected workspace, so you no longer hop between separate SharePoint forms and lists.",14,False,TEXT2,after=0)
cards=[("Request","Raise and approve a data purchase"),("Catalogue","See everything already purchased"),("Measure","Know which KPIs use which data"),("Source","Know your vendors and leverage"),("Govern","Track agreements and renewals")]
cw=Inches(2.28); y=Inches(4.0)
for i,(t1,t2) in enumerate(cards):
    xx=MX+i*(cw+Inches(0.1)); card(s,xx,y,cw,Inches(1.5))
    tt=tb(s,xx+Inches(0.18),y+Inches(0.24),cw-Inches(0.36),Inches(1.1)); par(tt,t1,14,True,INK,first=True,after=4); par(tt,t2,11.5,False,TEXT2,after=0)
footer(s,pg())
notes(s,"Plain-language definition. Emphasise the 'one place' benefit: today these are five disconnected SharePoint tools; Atlas joins them so you can see across them. Read the five verbs left to right - that's the lifecycle of any data asset, and it maps to the five modules in the left nav.")

# ---------- 3 One asset five tools
s=slide(); bg(s,SURFACE)
header(s,"The big idea","One asset, five tools - all connected")
flow=[("Data Requests","you raise a request"),("Data Catalogue","it becomes a record"),("KPI Catalogue","KPIs measure it"),("Vendor Hub","a vendor supplies it"),("TPA Hub","an agreement governs it")]
cw=Inches(2.18); y=Inches(2.3)
for i,(tl,ds) in enumerate(flow):
    xx=MX+i*(cw+Inches(0.16)); card(s,xx,y,cw,Inches(1.6))
    tt=tb(s,xx+Inches(0.16),y+Inches(0.22),cw-Inches(0.32),Inches(1.2)); par(tt,tl,13.5,True,INK,first=True,after=5); par(tt,ds,11.5,False,TEXT2,after=0)
    if i<4:
        ar=tb(s,xx+cw,y+Inches(0.5),Inches(0.16),Inches(0.5),anchor=MSO_ANCHOR.MIDDLE); par(ar,">",16,True,ORANGE,align=PP_ALIGN.CENTER,first=True,after=0)
t=tb(s,MX,Inches(4.3),CW,Inches(1.2))
par(t,"Because the five are connected, everything links to everything",16,True,INK,first=True,after=8)
bullets(t,[("From a request,","jump to the catalogue entry it creates."),("From an asset,","jump to its KPIs, its vendor and its agreement."),("From a KPI,","trace the exact asset and vendor behind the number."),("From an agreement,","see the asset and the strategic KPI that depend on it.")],size=13.5)
footer(s,pg())
notes(s,"The single most important concept: 'one asset, five tools.' Trace one example out loud. Then point out that the links are real and clickable in the tool - blue links and clickable rows jump you straight to the related record. This is what makes Atlas more than five separate lists.")

# ---------- 4 Screen anatomy
s=slide(); bg(s,SURFACE)
header(s,"How it is laid out","The screen, at a glance")
# wireframe
wx=MX; wy=Inches(2.0); ww=Inches(8.2); wh=Inches(4.6)
rect(s,wx,wy,ww,wh,fill=CARD,line=LINE,lw=1.0,shape=MSO_SHAPE.ROUNDED_RECTANGLE,radius=0.03)
# top bar
rect(s,wx,wy,ww,Inches(0.5),fill=INK,shape=MSO_SHAPE.ROUNDED_RECTANGLE,radius=0.04)
tt=tb(s,wx+Inches(0.15),wy,Inches(4),Inches(0.5),anchor=MSO_ANCHOR.MIDDLE); par(tt,"Suite   Cortex Atlas",10,True,WHITE,first=True,after=0)
rect(s,wx+ww-Inches(2.0),wy+Inches(0.12),Inches(1.0),Inches(0.26),fill=INK2,shape=MSO_SHAPE.ROUNDED_RECTANGLE,radius=0.4)
tt=tb(s,wx+ww-Inches(2.0),wy+Inches(0.12),Inches(1.0),Inches(0.26),anchor=MSO_ANCHOR.MIDDLE); par(tt,"Filters",8.5,True,MUTE,align=PP_ALIGN.CENTER,first=True,after=0)
# left nav
rect(s,wx,wy+Inches(0.5),Inches(1.9),wh-Inches(0.5),fill=INK2,shape=MSO_SHAPE.RECTANGLE)
nv=tb(s,wx+Inches(0.15),wy+Inches(0.65),Inches(1.7),Inches(3.6))
for i,(lab) in enumerate(["Command Centre","Data Requests","Data Catalogue","KPI Catalogue","Vendor Hub","TPA Hub","Vendor Hub Plus"]):
    par(nv,lab,9.5,i==0,RGBColor(0xC7,0xCC,0xD3) if i!=0 else ORANGE,first=(i==0),after=6)
# working surface cards
for i in range(3):
    rect(s,wx+Inches(2.1),wy+Inches(0.7)+i*Inches(1.25),ww-Inches(2.35),Inches(1.05),fill=SURFACE,line=LINE,lw=0.75,shape=MSO_SHAPE.ROUNDED_RECTANGLE,radius=0.06)
tt=tb(s,wx+Inches(2.25),wy+Inches(0.78),Inches(4),Inches(0.3)); par(tt,"Working surface - tiles, charts, tables",9.5,True,TEXT2,first=True,after=0)
# drawer
rect(s,wx+ww-Inches(2.5),wy+Inches(0.5),Inches(2.5),wh-Inches(0.5),fill=CARD,line=ORANGE,lw=1.25,shape=MSO_SHAPE.RECTANGLE)
tt=tb(s,wx+ww-Inches(2.4),wy+Inches(0.62),Inches(2.3),Inches(0.4)); par(tt,"Detail drawer",9.5,True,ORANGE_D,first=True,after=0)
# callouts
cx=MX+Inches(8.5)
calls=[("Top bar","module title, Live badge, and the Division / Financial-Year filters that apply everywhere."),
 ("Left navigation","grouped into Command Centre, The Estate (the 5 tools) and Intelligence."),
 ("Working surface","KPI tiles, charts and tables for the module you are in."),
 ("Detail drawer","slides in when you click a row - the full record, tabs and actions."),
 ("Ask Atlas","ask a question in plain English on the Command Centre.")]
cy=Inches(2.05)
for t1,t2 in calls:
    bt=tb(s,cx,cy,Inches(3.6),Inches(0.95)); par(bt,t1,12.5,True,INK,first=True,after=2); par(bt,t2,10.5,False,TEXT2,after=0); cy=cy+Inches(0.92)
footer(s,pg())
notes(s,"Teach the layout so nobody is lost. Four areas: the dark top bar (title + the Division and Financial-Year filters that apply across every module), the dark left nav (three groups), the light working surface (tiles/charts/tables), and the detail drawer that slides in from the right when you click a row. Ask Atlas lives on the Command Centre. Mention you can close the drawer with the X, by clicking outside it, or pressing Escape.")

# ---------- 5 Three things it does
s=slide(); bg(s,SURFACE)
header(s,"What it does for you","Three things, in every module")
layers=[("Keeps the record","The faithful as-is tool: forms, records, the approval workflow and the links between them. Your single source of truth.",BLUE),
 ("Does the busywork","Automation: auto-routing of requests, renewal pre-fill, write-to-catalogue, PII/PHI classification, and the 30/60/90-day reminders.",ORANGE),
 ("Gives you the answer","Intelligence: duplicate-buy alerts, quality flags, vendor value analytics, renewal prediction and the Ask Atlas assistant.",GREEN)]
cw=Inches(3.83); y=Inches(2.1)
for i,(tl,ds,col) in enumerate(layers):
    xx=MX+i*(cw+Inches(0.1)); card(s,xx,y,cw,Inches(4.3))
    rect(s,xx,y,cw,Inches(0.14),fill=col,shape=MSO_SHAPE.ROUNDED_RECTANGLE,radius=0.5)
    badge(s,xx+Inches(0.25),y+Inches(0.4),str(i+1),d=Inches(0.5),fill=col,fg=WHITE)
    tt=tb(s,xx+Inches(0.25),y+Inches(1.1),cw-Inches(0.5),Inches(3.0)); par(tt,tl,18,True,INK,first=True,after=8); par(tt,ds,13,False,TEXT2,after=0)
footer(s,pg())
notes(s,"Set expectations about how much the tool helps. It keeps your records (so you trust it), it does the repetitive work for you (so you save time and avoid mistakes), and it surfaces answers and flags (so you make better decisions). Reassure: every AI/analytics output is advisory - it shows a confidence read and a human always approves the action. Nothing happens automatically without you.")

# ---------- 6 Command Centre
s=slide(); bg(s,SURFACE)
header(s,"Module - Command Centre","Your landing page and estate overview")
twocol(s,"What it shows",[
 ("Headline numbers,","catalogued assets, contracted spend, in-flight requests, agreements expiring soon."),
 ("The connected estate,","the five tools as clickable tiles."),
 ("Cross-tool signals,","the handful of issues that need a decision, ranked by what is at stake."),
 ("Ask Atlas,","a question box that reads across all five tools.")],
 "What you can do",[
 ("Click any estate tile","to jump into that module."),
 ("Click a signal","to go straight to the record behind it."),
 ("Ask a question","by typing or using a suggested chip - you get a grounded answer with source chips."),
 ("Read the disclaimer","- figures are synthetic and illustrative.")])
footer(s,pg())
notes(s,"The Command Centre is home base. Tell users to start their day here: the signals tell you what needs attention (an agreement expiring behind a key KPI, a possible duplicate buy, a quality risk), and Ask Atlas answers cross-estate questions. Everything is a jumping-off point - clicking a signal or tile takes you to the underlying record.")

# ---------- 7 Data Requests - how the workflow works
s=slide(); bg(s,SURFACE)
header(s,"Module - Data Requests (DPRH)","How the approval workflow works")
# chain
steps=[("Submit","auto-status set"),("Admin review","routes it"),("Regional review","region-scoped"),("CDS committee","reg + global"),("Admin final","approves"),("Catalogue","as a draft")]
cw=Inches(1.95); y=Inches(2.15)
for i,(a,b) in enumerate(steps):
    xx=MX+i*(cw+Inches(0.04)); card(s,xx,y,cw,Inches(1.1))
    tt=tb(s,xx+Inches(0.12),y+Inches(0.16),cw-Inches(0.24),Inches(0.85)); par(tt,a,11.5,True,INK,first=True,after=2); par(tt,b,9.5,False,TEXT2,after=0)
    if i<5:
        ar=tb(s,xx+cw-Inches(0.02),y+Inches(0.35),Inches(0.12),Inches(0.4),anchor=MSO_ANCHOR.MIDDLE); par(ar,">",12,True,ORANGE,align=PP_ALIGN.CENTER,first=True,after=0)
t=tb(s,MX,Inches(3.5),CW,Inches(3.0))
par(t,"Two things decide where a request goes - automatically, on submit:",14.5,True,INK,first=True,after=8)
bullets(t,[
 ("National vs Sub-national.","Sub-national requests go to offline review with local teams - they skip the committee chain."),
 ("Cost vs threshold.","For national requests, if the anticipated cost is above your division's threshold it becomes In data sourcing pre review; if below, Submitted for CILT review."),
 ("Roles see their own tab.","Requester fills the form; Admin, Regional and Global reviewers each see and act on their own section, which opens only when the request reaches that step."),
 ("Statuses are the language.","Each request always shows exactly where it is - e.g. Submitted for validation, In CILT CDSP review, Recommendation available, Approved in contracting.")],size=12.5)
footer(s,pg())
notes(s,"This is the heart of the tool, so spend time here. Walk the chain left to right. Then the two routing rules - national/sub-national and cost-vs-threshold - which the tool applies automatically when you submit. Explain the four roles and that tabs open progressively (you won't see the committee tab until it gets there). The status names are the shared vocabulary - the training guide lists them all with meanings.")

# ---------- 8 Data Requests - what you can do
s=slide(); bg(s,SURFACE)
header(s,"Module - Data Requests (DPRH)","What you can do")
twocol(s,"Everyday tasks",[
 ("Raise a request","with New data request - the form shows the routing and flags duplicates as you type."),
 ("Renew","by choosing Continuing and the data being renewed - the form pre-fills from the catalogue."),
 ("Review and act","by switching your role and using the action buttons (route, recommend, approve, reject)."),
 ("Move to catalogue","once approved - it creates a draft catalogue entry to complete.")],
 "Helpful to know",[
 ("Duplicate alerts","appear if your request overlaps something already owned - review before spending."),
 ("Sub-national","requests are handled offline by local teams."),
 ("Every field is captured","across the Request, Scope, and review tabs - nothing is re-entered later."),
 ("Actions are role-based","- if you see 'no actions for this role', switch to the role that owns this step.")])
footer(s,pg())
notes(s,"Now the action view. Demo live if you can: click New data request, fill a couple of fields and watch the routing/duplicate panel update; submit; then switch roles (Requester -> Admin -> Regional -> Global -> Admin) and click the action buttons to walk it to Approved; then Move to Data Catalogue. Stress the role switch: the action buttons change with your role and the request's current status.")

# ---------- 9 Data Catalogue
s=slide(); bg(s,SURFACE)
header(s,"Module - Data Catalogue","The record of everything purchased")
twocol(s,"What it shows",[
 ("Every purchased asset,","with its vendor, therapeutic area and brand."),
 ("A quality tier,","Gold, Silver or Bronze, per asset."),
 ("Sensitivity flags,","PII and PHI."),
 ("A full record,","identity, classification, coverage, commercials and access - plus its lineage."),
 ("Quality flags,","when a low-tier asset feeds a strategic KPI.")],
 "What you can do",[
 ("Search before you buy","- check whether the data already exists."),
 ("Add a manual entry","- the tool auto-classifies PII/PHI and suggests a quality tier as you fill it in."),
 ("Complete and publish","a draft that arrived from an approved request."),
 ("Jump to related records","- the vendor, the KPIs it feeds, its agreement.")])
footer(s,pg())
notes(s,"The catalogue is the single source of truth and the answer to 'do we already own this?'. Show opening an asset record and scrolling the grouped fields. Then New manual entry - pick a patient-level claims type and watch PHI flip to Yes with a suggested Gold tier. Mention drafts: when a request is approved and moved here, it lands as a Draft to be completed and published.")

# ---------- 10 KPI Catalogue
s=slide(); bg(s,SURFACE)
header(s,"Module - KPI Catalogue","The metrics built on the data")
twocol(s,"What it shows",[
 ("Every KPI and metric,","with definition, methodology, cadence and contacts."),
 ("Scope,","Global, Regional or Local."),
 ("Lineage,","the source asset (and therefore vendor) behind each number."),
 ("Quality risk,","a flag when a strategic KPI rests on a low-tier asset.")],
 "What you can do",[
 ("Trace any number","to its source in two clicks."),
 ("Add a KPI or metric","and link its source asset - lineage and quality risk compute automatically."),
 ("Judge how much to trust a metric","by seeing the quality of the data beneath it.")])
footer(s,pg())
notes(s,"The KPI catalogue answers 'where does this number come from and how good is it?'. Show the lineage chain in a KPI record. Then Add KPI / metric - pick a Bronze source asset with Global scope and the tool flags the quality risk live. This turns institutional knowledge that used to live in people's heads into something anyone can look up.")

# ---------- 11 Vendor Hub
s=slide(); bg(s,SURFACE)
header(s,"Module - Vendor Hub","Who supplies your data")
twocol(s,"What it shows",[
 ("Every vendor,","what data they hold and how to access it."),
 ("Purchased vs available,","what you buy versus everything they offer."),
 ("Contacts and guidance,","per vendor."),
 ("A link up to the value layer","- Vendor Hub Plus.")],
 "What you can do",[
 ("Look up a vendor","and what we license from them."),
 ("Spot untapped coverage","- offerings you do not yet buy."),
 ("Open Vendor Hub Plus","for the value and risk analytics.")])
footer(s,pg())
notes(s,"Vendor Hub is the inventory - who our vendors are and what we buy from each. Keep it brief; it sets up Vendor Hub Plus, which is where the analytics live. Point out 'purchased vs available' - it is the basis for negotiation leverage on the next slide.")

# ---------- 12 Vendor Hub Plus
s=slide(); bg(s,SURFACE)
header(s,"Module - Vendor Hub Plus","The vendor value and risk layer")
plus=[("Quality score","each vendor's Gold/Silver/Bronze mix rolled up to a grade"),
 ("Spend concentration","each vendor's share of spend, with an estate-wide HHI"),
 ("Irreplaceability","what breaks if a vendor leaves - assets and KPIs at stake"),
 ("Cost-per-quality","whether we are overpaying for low-tier data"),
 ("Leverage","purchased vs available, as a negotiation lever")]
cw=Inches(2.3); y=Inches(2.2)
for i,(tl,ds) in enumerate(plus):
    xx=MX+i*(cw+Inches(0.1)); card(s,xx,y,cw,Inches(3.1))
    badge(s,xx+Inches(0.2),y+Inches(0.24),str(i+1),d=Inches(0.44))
    tt=tb(s,xx+Inches(0.2),y+Inches(0.9),cw-Inches(0.4),Inches(2.0)); par(tt,tl,13,True,INK,first=True,after=6); par(tt,ds,11,False,TEXT2,after=0)
t=tb(s,MX,Inches(5.6),CW,Inches(0.7))
par(t,"All five recompute automatically as the catalogue changes - and all are advisory; a human makes the call.",13,True,INK,first=True,after=0)
footer(s,pg())
notes(s,"This is the standout analytics view. Explain each of the five in user terms - who delivers quality, where spend is concentrated (HHI: higher = more concentrated, more risk), who we cannot easily replace, where we overpay, and where we have room to negotiate. Stress these update automatically as the catalogue changes, and they are advisory inputs to a renewal decision, not automatic actions.")

# ---------- 13 TPA Hub
s=slide(); bg(s,SURFACE)
header(s,"Module - TPA Hub","Agreements and renewal reminders")
twocol(s,"What it shows",[
 ("Every agreement,","scope, signatories, dates and owner."),
 ("An expiry runway,","what is due in 30, 60 and 90 days."),
 ("Reminder status,","which 30/60/90 reminders have fired."),
 ("Cross-tool warnings,","when an expiring agreement sits behind a strategic KPI."),
 ("A renewal recommendation","per agreement.")],
 "What you can do",[
 ("See what is expiring","and act before leverage is lost."),
 ("Open an agreement","for its full detail and reminder timeline."),
 ("Create a renewal request","- raises a pre-filled request in Data Requests."),
 ("Follow the link","to the asset and KPI that depend on it.")])
footer(s,pg())
notes(s,"TPA Hub is both a register and a live reminder engine. The runway view (30/60/90) is what most users will check. The key upgrade over a plain reminder: it tells you the impact - 'this expires in 21 days and it is the only agreement behind a Tier-1 KPI'. And Create renewal request links straight back into the request workflow, so renewing is one click to start.")

# ---------- 14 Ask Atlas + signals
s=slide(); bg(s,SURFACE)
header(s,"The assistant","Ask Atlas and the cross-tool signals")
twocol(s,"Ask Atlas (Command Centre)",[
 ("Type a question","in plain English, or tap a suggested one."),
 ("Get a grounded answer","that reads across all five tools."),
 ("See the sources","- the feeds the answer is based on."),
 ("Follow the links","in the answer to the records.")],
 "How to trust it",[
 ("Confidence read","- every answer shows how confident it is."),
 ("Advisory only","- it informs you; it never acts on its own."),
 ("A human approves","any action that follows."),
 ("Grounded, not invented","- it cites the source rows.")])
footer(s,pg())
notes(s,"Demystify the AI. It is an assistant that reads across the connected data and answers questions or pushes signals - it does not make decisions. Show one Ask Atlas question. Then point at the confidence line and the 'advisory only - a human approves' cue that appears on every AI output. This is important for a regulated environment: nothing is automated without a person.")

# ---------- 15 Good to know
s=slide(); bg(s,SURFACE)
header(s,"Good to know","A few things that apply everywhere")
items=[("Filters apply across modules","The Division and Financial-Year selectors in the top bar filter the tables and charts in every module."),
 ("Everything is linked","Blue links and clickable rows jump you to the related record; the drawer closes when you navigate."),
 ("Close the drawer easily","Use the X, click the dark area outside it, or press Escape."),
 ("AI is advisory","Insights show a confidence read; a human always approves the action."),
 ("Synthetic data","All names and figures are illustrative for training - not live client data."),
 ("Need detail?","The User Training Guide has step-by-step instructions for every task.")]
cw=Inches(5.85)
for i,(t1,t2) in enumerate(items):
    col=i%2; rowi=i//2; xx=MX+col*(cw+Inches(0.3)); yy=Inches(2.05)+rowi*Inches(1.5)
    card(s,xx,yy,cw,Inches(1.35))
    tt=tb(s,xx+Inches(0.25),yy+Inches(0.2),cw-Inches(0.5),Inches(1.0)); par(tt,t1,13.5,True,INK,first=True,after=3); par(tt,t2,11.5,False,TEXT2,after=0)
footer(s,pg())
notes(s,"Wrap-up of cross-cutting behaviours people ask about: the global filters, the clickable links, how to close the drawer, the advisory nature of AI, and the synthetic-data note. Send them to the Training Guide for how-to detail.")

# ---------- 16 Quick start
s=slide(); bg(s,INK)
rect(s,0,0,SW,Inches(0.16),fill=ORANGE)
t=tb(s,MX+Inches(0.45),Inches(0.7),Inches(10),Inches(0.4)); par(t,"QUICK START",13,True,ORANGE,first=True,after=0)
t=tb(s,MX,Inches(1.1),Inches(11.8),Inches(0.8)); par(t,"Try these five things",34,True,WHITE,first=True,after=0)
steps=[("1","Open the Command Centre","read the signals and ask Atlas a question."),
 ("2","Raise a request","Data Requests > New data request; watch the routing and duplicate check."),
 ("3","Walk it through","open a request, switch roles, and use the action buttons to approve it."),
 ("4","Move it to the catalogue","then Complete & publish the draft."),
 ("5","Explore the intelligence","Vendor Hub Plus, KPI lineage and the TPA runway.")]
y=Inches(2.2)
for n,tl,ds in steps:
    badge(s,MX,y,n,d=Inches(0.5),fill=ORANGE,fg=INK)
    t=tb(s,MX+Inches(0.8),y-Inches(0.02),Inches(11.0),Inches(0.9)); par(t,tl,16.5,True,WHITE,first=True,after=2); par(t,ds,12.5,False,RGBColor(0xC7,0xCC,0xD3),after=0); y=y+Inches(0.92)
t=tb(s,MX,Inches(6.9),Inches(11.8),Inches(0.4)); par(t,"For step-by-step instructions on every feature, see the Cortex Atlas User Training Guide.",11.5,False,MUTE,first=True,after=0)
notes(s,"End with five concrete things to try so users leave able to do something immediately. Point them to the User Training Guide for the full how-to. Encourage them to click around - it is safe; the prototype uses synthetic data and resets on refresh.")

prs.save("/home/user/datasharing/Cortex-Atlas-Explainer.pptx")
print("Explainer deck saved with", len(prs.slides._sldIdLst), "slides")
